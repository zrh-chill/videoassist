from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "videoassist_bilibili_v1_prd.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, after=5, before=0, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    format_paragraph(p)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        format_paragraph(p, after=3)
        set_run_font(p.add_run(item))


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        format_paragraph(p, after=3)
        set_run_font(p.add_run(item))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    size = 16 if level == 1 else 12
    set_run_font(p.add_run(text), size=size, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr()
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "1F4E78")
        set_cell_borders(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, after=0, line=1.15)
        set_run_font(p.add_run(header), size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[idx])
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            set_cell_borders(cell)
            set_cell_margin(cell)
            if row_idx % 2:
                set_cell_shading(cell, "F3F7FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if alignments and alignments[idx] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_paragraph(p, after=0, line=1.15)
            set_run_font(p.add_run(str(value)), size=9.5)
            if widths:
                cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("视频转写与总结应用产品需求文档"), size=22, bold=True)

    subtitle = doc.add_paragraph()
    format_paragraph(subtitle, after=14)
    set_run_font(subtitle.add_run("首版目标平台  B站"), size=12, bold=True, color="404040")

    add_table(
        doc,
        ["文档项", "内容"],
        [
            ["产品定位", "个人使用的视频转写与 AI 总结工具"],
            ["首版范围", "本地视频与 B 站视频链接，支持 B 站 UP 主追踪"],
            ["核心目标", "完成视频获取、音频提取、全文转写、AI 总结和结果沉淀的闭环"],
            ["数据存储", "关系型数据库为主，并支持导出 Excel"],
        ],
        widths=[1.25, 5.45],
        alignments=["center", "left"],
    )

    add_heading(doc, "1 产品概述")
    add_body(
        doc,
        "本产品面向个人用户，用于将本地视频或 B 站视频快速处理为完整文稿和结构化总结。系统统一记录视频来源、处理过程、转写内容、总结结果和异常信息，方便后续查阅、筛选与导出。",
    )
    add_body(
        doc,
        "首版重点：跑通单条视频处理和 UP 主最新视频批量处理两个核心场景，以稳定、可追踪和可重新执行为主要要求。",
        bold_lead="首版重点：",
    )

    add_heading(doc, "2 产品目标与范围")
    add_heading(doc, "2.1 产品目标", level=2)
    add_bullets(
        doc,
        [
            "支持上传本地视频或输入 B 站视频链接创建处理任务。",
            "使用 FFmpeg 从视频中提取音频，并通过语音转文字模型生成完整文稿。",
            "使用语言模型根据系统提示词生成结构化 AI 总结。",
            "支持添加 B 站 UP 主，并获取其最新 N 条视频批量处理。",
            "以一条视频一条主记录的方式展示处理状态、结果和过程信息。",
        ],
    )
    add_heading(doc, "2.2 首版不包含", level=2)
    add_bullets(
        doc,
        [
            "除 B 站之外的视频平台适配。",
            "多人账户、团队协作、权限管理与内容分享。",
            "人工校对工作台、字幕编辑和视频剪辑。",
            "复杂知识库、向量检索和基于全部视频的问答。",
            "移动端原生应用。",
        ],
    )

    add_heading(doc, "3 核心使用流程")
    add_numbered(
        doc,
        [
            "用户上传本地视频、粘贴 B 站视频链接，或从已追踪 UP 主中发现新视频。",
            "系统建立视频记录并检查是否已经处理，避免同一 B 站视频重复入库。",
            "系统获取视频文件，调用 FFmpeg 提取统一格式的音频。",
            "系统调用语音转文字模型，生成完整文稿并保存模型与耗时信息。",
            "系统调用语言模型，按当前总结提示词生成结构化总结。",
            "用户在视频列表和详情页查看结果，必要时重新转写或重新总结。",
            "用户可筛选数据，并将视频记录与主要结果导出为 Excel。",
        ],
    )

    add_heading(doc, "4 功能需求")
    add_table(
        doc,
        ["模块", "首版功能", "说明"],
        [
            ["视频任务", "本地上传", "选择本地视频文件后创建任务，记录文件名、大小和创建时间。"],
            ["视频任务", "B站链接导入", "识别常见 B 站视频链接，读取视频基础信息并创建任务。"],
            ["视频获取", "下载与去重", "按 BVID 或规范化链接去重；下载失败时保留原因并支持重试。"],
            ["音频处理", "FFmpeg 提取", "将视频转换为语音模型可接受的统一音频格式，并记录执行状态。"],
            ["全文转写", "S2T 转写", "生成完整文稿；建议同时保留带起止时间的分段结果。"],
            ["AI总结", "结构化总结", "输出一句话摘要、核心要点、详细总结和关键词。"],
            ["结果管理", "列表与详情", "按视频展示来源、状态、标题、UP主、发布时间、文稿和总结。"],
            ["结果管理", "重新处理", "允许从下载、转写或总结阶段重新执行，旧结果可被最新成功结果替代。"],
            ["数据导出", "Excel 导出", "导出筛选后的主要视频信息、完整文稿和总结字段。"],
        ],
        widths=[1.05, 1.35, 4.3],
        alignments=["center", "center", "left"],
    )

    add_heading(doc, "4.1 B站 UP 主追踪", level=2)
    add_bullets(
        doc,
        [
            "添加 UP 主：填写 UP 主主页链接或 UID，并保存昵称等基础信息。",
            "追踪设置：配置每次检查最新 N 条视频，首版建议默认 N 为 5。",
            "手动检查：点击检查更新后拉取最新视频，与已有 BVID 去重。",
            "自动入队：发现新视频后自动创建处理任务，可在追踪设置中开启或关闭。",
            "追踪结果：记录上次检查时间、新增数量和失败原因。",
        ],
    )

    add_heading(doc, "4.2 AI 总结与系统配置", level=2)
    add_bullets(
        doc,
        [
            "配置语音转文字模型、接口地址、密钥引用和基础参数。",
            "配置语言模型、接口地址、密钥引用和基础参数。",
            "编辑总结系统提示词，并保存当前使用的提示词版本。",
            "当文稿超过模型上下文限制时，系统应支持分段总结后再汇总。",
            "重新总结时使用当前配置，同时在结果中记录实际模型与提示词版本。",
        ],
    )

    add_heading(doc, "5 页面需求")
    add_table(
        doc,
        ["页面", "主要内容", "主要操作"],
        [
            ["视频列表", "状态、标题、来源、UP主、发布时间、创建时间、错误摘要", "搜索、筛选、查看详情、重试、导出"],
            ["添加视频", "本地上传区、B站链接输入框", "上传文件、解析链接、创建任务"],
            ["视频详情", "基础信息、处理进度、完整文稿、AI总结、过程日志", "复制、重新转写、重新总结、重试"],
            ["UP主追踪", "UP主列表、最新 N 条、自动处理开关、上次检查结果", "添加、编辑、删除、立即检查"],
            ["系统设置", "FFmpeg、S2T、语言模型、总结提示词和存储配置", "保存配置、测试连接"],
        ],
        widths=[1.1, 3.45, 2.15],
        alignments=["center", "left", "left"],
    )

    add_heading(doc, "6 处理状态与异常")
    add_body(doc, "视频主记录应显示一个明确的当前状态，并保存各阶段的开始时间、结束时间和错误信息。")
    add_table(
        doc,
        ["状态", "含义", "允许操作"],
        [
            ["等待处理", "任务已创建，等待执行", "取消或开始处理"],
            ["获取视频中", "正在读取本地文件或下载 B 站视频", "查看进度"],
            ["提取音频中", "正在执行 FFmpeg", "查看日志"],
            ["转写中", "正在调用 S2T 模型", "查看进度"],
            ["总结中", "正在调用语言模型", "查看进度"],
            ["已完成", "文稿和总结均已生成", "查看、导出、重新处理"],
            ["处理失败", "任一阶段执行失败", "查看原因并从失败阶段重试"],
        ],
        widths=[1.25, 3.35, 2.1],
        alignments=["center", "left", "left"],
    )

    add_heading(doc, "7 数据记录")
    add_body(doc, "数据库作为主存储，Excel 作为导出格式。一条视频对应一条视频主记录，转写、总结和执行日志可作为关联记录保存。")
    add_table(
        doc,
        ["数据对象", "主要字段"],
        [
            ["视频", "来源类型、BVID、原始链接、本地路径、标题、UP主、发布时间、时长、当前状态、错误信息"],
            ["转写", "视频ID、完整文稿、分段文本、S2T模型、语言、处理耗时、创建时间"],
            ["总结", "视频ID、一句话摘要、核心要点、详细总结、关键词、语言模型、提示词版本、创建时间"],
            ["UP主", "UID、主页链接、昵称、最新N条、自动处理、上次检查时间、检查结果"],
            ["执行日志", "视频ID、处理阶段、开始时间、结束时间、状态、错误详情"],
        ],
        widths=[1.25, 5.45],
        alignments=["center", "left"],
    )

    add_heading(doc, "8 非功能要求")
    add_bullets(
        doc,
        [
            "个人使用优先，首版可不提供登录功能，但应用默认只在受信任环境运行。",
            "访问密钥不得直接展示在列表、日志或导出文件中。",
            "单个阶段失败不应丢失已完成的数据，重新执行应能从指定阶段开始。",
            "页面刷新或应用重启后，任务状态和历史结果仍可恢复。",
            "B站内容获取能力应遵守平台规则，登录态、限流或链接失效时给出可理解的提示。",
        ],
    )

    add_heading(doc, "9 首版验收标准")
    add_bullets(
        doc,
        [
            "能够上传一个本地视频并生成完整文稿与结构化总结。",
            "能够输入一个有效 B 站视频链接并完成下载、转写和总结。",
            "相同 BVID 重复导入时不会产生重复视频记录。",
            "能够添加一个 B 站 UP 主，获取其最新 N 条视频并为新增视频创建任务。",
            "能够在列表中查看全部视频的当前处理状态，并在详情中查看完整文稿、总结和错误日志。",
            "转写或总结失败后可以重试，并在成功后正确更新结果。",
            "能够将筛选后的视频记录、完整文稿和总结导出为 Excel。",
            "能够修改总结提示词并使用新配置重新生成总结。",
        ],
    )

    add_heading(doc, "10 建议实施顺序")
    add_numbered(
        doc,
        [
            "完成视频任务、数据库和处理状态基础能力。",
            "接入本地视频、FFmpeg、S2T 和语言模型，跑通单条视频闭环。",
            "接入 B 站视频解析、下载和 BVID 去重。",
            "实现视频列表、详情页、失败重试和 Excel 导出。",
            "实现 UP 主追踪、最新 N 条检查和自动创建任务。",
            "补充系统配置、提示词版本和长文稿分段总结。",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
